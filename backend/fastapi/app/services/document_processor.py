from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


class DocumentProcessingError(Exception):
    """Raised when a document cannot be processed."""


class DocumentProcessor:
    """
    Handles document text extraction, cleaning, and chunking.

    Supported formats:
    - PDF
    - DOCX
    - TXT
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".txt",
    }

    DEFAULT_CHUNK_SIZE = 1000
    DEFAULT_CHUNK_OVERLAP = 150

    @classmethod
    def extract_text(
        cls,
        filename: str,
        content: bytes,
    ) -> str:
        """Extract text from a supported document."""

        extension = Path(filename).suffix.lower()

        if extension not in cls.SUPPORTED_EXTENSIONS:
            raise DocumentProcessingError(
                f"Unsupported file type: {extension or 'unknown'}. "
                "Supported types: PDF, DOCX, TXT."
            )

        if not content:
            raise DocumentProcessingError(
                "The uploaded file is empty."
            )

        try:
            if extension == ".pdf":
                text = cls._extract_pdf(content)

            elif extension == ".docx":
                text = cls._extract_docx(content)

            elif extension == ".txt":
                text = cls._extract_txt(content)

            else:
                raise DocumentProcessingError(
                    f"Unsupported file type: {extension}"
                )

        except DocumentProcessingError:
            raise

        except Exception as exc:
            raise DocumentProcessingError(
                f"Failed to extract text from {filename}: {exc}"
            ) from exc

        cleaned_text = cls.clean_text(text)

        if not cleaned_text:
            raise DocumentProcessingError(
                "No readable text was found in the document."
            )

        return cleaned_text

    @staticmethod
    def _extract_pdf(content: bytes) -> str:
        """Extract text from all PDF pages."""

        reader = PdfReader(BytesIO(content))

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            try:
                page_text = page.extract_text() or ""

            except Exception as exc:
                raise DocumentProcessingError(
                    f"Could not extract text from PDF "
                    f"page {page_number}: {exc}"
                ) from exc

            if page_text.strip():
                pages.append(page_text)

        return "\n\n".join(pages)

    @staticmethod
    def _extract_docx(content: bytes) -> str:
        """Extract paragraphs and table contents from DOCX."""

        document = DocxDocument(BytesIO(content))

        parts = []

        # Extract normal paragraphs.
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Extract table contents.
        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)

    @staticmethod
    def _extract_txt(content: bytes) -> str:
        """Decode a plain text file."""

        try:
            return content.decode("utf-8")

        except UnicodeDecodeError:
            try:
                return content.decode("utf-8-sig")

            except UnicodeDecodeError:
                return content.decode("latin-1")

    @staticmethod
    def clean_text(text: str) -> str:
        """Normalize extracted document text."""

        if not text:
            return ""

        # Remove null characters.
        text = text.replace("\x00", " ")

        # Normalize line endings.
        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")

        # Collapse repeated spaces and tabs.
        text = re.sub(r"[ \t]+", " ", text)

        # Remove excessive blank lines.
        text = re.sub(r"\n{3,}", "\n\n", text)

        # Remove whitespace from individual lines.
        lines = [
            line.strip()
            for line in text.split("\n")
        ]

        text = "\n".join(lines)

        return text.strip()

    @classmethod
    def chunk_text(
        cls,
        text: str,
        chunk_size: int | None = None,
        chunk_overlap: int | None = None,
    ) -> list[str]:
        """
        Split cleaned text into overlapping chunks.

        The processor first tries to preserve paragraph boundaries.
        Long paragraphs are split by words.
        """

        if not text:
            return []

        chunk_size = (
            chunk_size
            if chunk_size is not None
            else cls.DEFAULT_CHUNK_SIZE
        )

        chunk_overlap = (
            chunk_overlap
            if chunk_overlap is not None
            else cls.DEFAULT_CHUNK_OVERLAP
        )

        if chunk_size <= 0:
            raise ValueError(
                "chunk_size must be greater than zero."
            )

        if chunk_overlap < 0:
            raise ValueError(
                "chunk_overlap cannot be negative."
            )

        if chunk_overlap >= chunk_size:
            raise ValueError(
                "chunk_overlap must be smaller than chunk_size."
            )

        paragraphs = [
            paragraph.strip()
            for paragraph in re.split(
                r"\n\s*\n",
                text,
            )
            if paragraph.strip()
        ]

        chunks: list[str] = []
        current = ""

        for paragraph in paragraphs:

            # Normal-sized paragraph.
            if len(paragraph) <= chunk_size:

                candidate = (
                    f"{current}\n\n{paragraph}"
                    if current
                    else paragraph
                )

                if len(candidate) <= chunk_size:
                    current = candidate
                    continue

                if current:
                    chunks.append(current)

                overlap_text = (
                    current[-chunk_overlap:]
                    if chunk_overlap > 0
                    else ""
                )

                current = (
                    f"{overlap_text}\n\n{paragraph}"
                ).strip()

                if len(current) > chunk_size:
                    chunks.extend(
                        cls._split_long_text(
                            current,
                            chunk_size,
                            chunk_overlap,
                        )
                    )
                    current = ""

            # Very large paragraph.
            else:
                if current:
                    chunks.append(current)
                    current = ""

                chunks.extend(
                    cls._split_long_text(
                        paragraph,
                        chunk_size,
                        chunk_overlap,
                    )
                )

        if current:
            chunks.append(current)

        return [
            chunk.strip()
            for chunk in chunks
            if chunk.strip()
        ]

    @staticmethod
    def _split_long_text(
        text: str,
        chunk_size: int,
        chunk_overlap: int,
    ) -> list[str]:
        """Split long text by words while preserving overlap."""

        words = text.split()

        if not words:
            return []

        chunks: list[str] = []
        current_words: list[str] = []
        current_length = 0

        for word in words:

            additional_length = (
                len(word)
                if not current_words
                else len(word) + 1
            )

            if (
                current_words
                and current_length + additional_length
                > chunk_size
            ):
                chunks.append(
                    " ".join(current_words)
                )

                # Preserve the last part of the previous
                # chunk as overlap.
                overlap_words: list[str] = []
                overlap_length = 0

                for previous_word in reversed(
                    current_words
                ):
                    extra = (
                        len(previous_word)
                        if not overlap_words
                        else len(previous_word) + 1
                    )

                    if (
                        overlap_length + extra
                        > chunk_overlap
                    ):
                        break

                    overlap_words.insert(
                        0,
                        previous_word,
                    )

                    overlap_length += extra

                current_words = overlap_words
                current_length = overlap_length

            current_words.append(word)
            current_length = len(
                " ".join(current_words)
            )

        if current_words:
            chunks.append(
                " ".join(current_words)
            )

        return chunks